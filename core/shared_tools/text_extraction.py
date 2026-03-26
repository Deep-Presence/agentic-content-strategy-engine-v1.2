"""Shared text extraction from document files.

Supports .md, .txt, .pdf (via pdfplumber), and .docx (via python-docx).
Used by both the knowledge doc upload service and the s1 pipeline step.

Two entry points:
  - ``extract_text(file_path)`` — reads from a local ``Path`` (legacy).
  - ``extract_text_from_bytes(content, suffix)`` — reads from in-memory bytes
    via ``io.BytesIO``.  Added in Phase 6 (R2 migration) so callers can
    extract text from StorageBackend content without touching the filesystem.
"""
from __future__ import annotations

import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)


def extract_text(file_path: Path) -> str:
    """Extract plain text from a file based on its extension.

    Returns empty string on failure or unsupported file type.
    """
    suffix = file_path.suffix.lower()

    if suffix in (".md", ".txt"):
        return file_path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".pdf":
        try:
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                return "\n\n".join(
                    page.extract_text() or "" for page in pdf.pages
                )
        except Exception as e:
            logger.warning("Failed to extract text from PDF %s: %s", file_path, e)
            return ""

    if suffix == ".docx":
        try:
            from docx import Document

            doc = Document(str(file_path))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            logger.warning("Failed to extract text from DOCX %s: %s", file_path, e)
            return ""

    return ""


def extract_text_from_bytes(content: bytes, suffix: str) -> str:
    """Extract plain text from in-memory bytes based on file extension.

    Works identically to :func:`extract_text` but accepts raw bytes and a
    file extension instead of a filesystem ``Path``.  Uses ``io.BytesIO``
    for PDF/DOCX so no temporary files are created.

    Returns empty string on failure or unsupported file type.
    """
    if not content:
        return ""

    suffix = suffix.lower()

    if suffix in (".md", ".txt"):
        return content.decode("utf-8", errors="replace")

    if suffix == ".pdf":
        try:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(content)) as pdf:
                return "\n\n".join(
                    page.extract_text() or "" for page in pdf.pages
                )
        except Exception as e:
            logger.warning("Failed to extract text from PDF bytes: %s", e)
            return ""

    if suffix == ".docx":
        try:
            from docx import Document

            doc = Document(io.BytesIO(content))
            return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            logger.warning("Failed to extract text from DOCX bytes: %s", e)
            return ""

    return ""
